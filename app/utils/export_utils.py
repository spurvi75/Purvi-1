"""
Excel export utility for the Reports module, using openpyxl.
"""

import io
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def export_bookings_to_excel(bookings, title="Booking Report"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    headers = [
        "Booking No.", "Employee Name", "Employee Email", "Mobile", "Grade", "Gender",
        "City", "Guest House", "Room No.", "Check-in", "Check-out",
        "Status", "HOD Approval Status", "HOD Email", "Purpose", "Created Date",
    ]

    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)

    ws.append([])
    header_row_idx = 3
    ws.append(headers)
    header_fill = PatternFill(start_color="0D2C54", end_color="0D2C54", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for b in bookings:
        ws.append([
            b.booking_number, b.employee_name, b.employee_email, b.mobile_number,
            b.grade, b.gender, b.city, b.guest_house_name, b.room_number,
            b.check_in_date.strftime("%Y-%m-%d") if isinstance(b.check_in_date, date) else b.check_in_date,
            b.check_out_date.strftime("%Y-%m-%d") if isinstance(b.check_out_date, date) else b.check_out_date,
            b.booking_status, b.hod_approval_status, b.hod_email, b.purpose_of_visit,
            b.created_date.strftime("%Y-%m-%d %H:%M") if b.created_date else "",
        ])

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=header_row_idx, column=col_idx).column_letter].width = max(14, len(header) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
